from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("outputs/kirolos-negative-actions-doc/kirolos-medhat-actions-2026-04-28.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text_direction(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    text_direction = OxmlElement("w:textDirection")
    text_direction.set(qn("w:val"), "tbRl")


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_rtl(p)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, 16, True, "143A5A")
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
    else:
        set_run_font(run, 13, True, "1F5B7A")
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        set_paragraph_rtl(p)
        p.paragraph_format.right_indent = Cm(0.35)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"- {item}")
        set_run_font(run, 10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, head in enumerate(headers):
        hdr[idx].text = ""
        set_cell_shading(hdr[idx], "143A5A")
        p = hdr[idx].paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(head)
        set_run_font(run, 10, True, "FFFFFF")
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            hdr[idx].width = widths[idx]

    for row in rows:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            set_paragraph_rtl(p)
            run = p.add_run(str(val))
            set_run_font(run, 9.5)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[idx].width = widths[idx]
    doc.add_paragraph()
    return table


def add_copy_block(doc, title, lines):
    add_heading(doc, title, 2)
    p = doc.add_paragraph()
    set_paragraph_rtl(p)
    for i, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, 10.5)
        if i < len(lines) - 1:
            run.add_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.4)
section.right_margin = Cm(1.4)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), "Arial")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
set_paragraph_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
r = title.add_run("تعديلات Google Ads الجاهزة للتنفيذ")
set_run_font(r, 20, True, "143A5A")

subtitle = doc.add_paragraph()
set_paragraph_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
r = subtitle.add_run("Kirolos Medhat | الفترة: 14 أبريل 2026 - 28 أبريل 2026")
set_run_font(r, 11, False, "5B6770")

add_heading(doc, "الخلاصة التنفيذية", 1)
add_bullets(
    doc,
    [
        "حوّلي كامبين اورام الثدي - الغده - الكبد إلى Maximize Conversions.",
        "لا تحوّلي كامبين بنكرياس - قولون - مراره إلى Max Conversions دلوقتي.",
        "ماتعمليش Portfolio/shared budget.",
        "الـ Location يفضل مصر كلها مع Presence only.",
        "الـ Schedule يبدأ من 12 ظهرًا إلى 12 منتصف الليل.",
        "الأولوية: negatives ثم keywords ثم تحسين الإعلانات ثم ضبط الميزانية.",
    ],
)

add_heading(doc, "Negative Keywords - Account Level", 1)
account_negatives = [
    "حميد", "الحميد", "الحميدة", "اعراض", "أعراض", "اسباب", "أسباب",
    "ما هو", "ما هي", "هل", "متى", "كم", "مدة", "تجربتي", "تجربة",
    "فيديو", "بالليزر", "كيماوي", "الكيماوي", "العلاج الكيماوي",
    "العلاج الهرموني", "العلاج الموجه", "بهية", "بهيه", "مركز بهية",
    "مركز بهيه", "سعر جلسة", "نسبة نجاح", "بعد العملية", "ماذا يحدث", "خطير",
]
add_copy_block(doc, "ضيفي Phrase Match:", account_negatives)

competitors = [
    "اشرف الزيات", "أشرف الزيات", "تامر النحاس", "د تامر النحاس",
    "دكتور تامر النحاس", "خالد نجيب", "د خالد نجيب", "كريم مشهور",
    "محمد العشري", "احمد الجويلي", "أحمد الجويلي", "احمد شكري",
    "أحمد شكري", "حسام عمر", "رامي غالي", "شريف نجيب", "لؤي قاسم",
    "محمد بسيوني", "محمد عصمت", "محمد متولي", "محمد محمود عبد الحكيم",
    "نهاد مكاوى", "هاني وليم",
]
add_copy_block(doc, "Competitor Negatives - Phrase Match:", competitors)

add_heading(doc, "Negative Keywords حسب الكامبين", 1)
add_table(
    doc,
    ["الكامبين", "الكلمات السلبية", "ملاحظات"],
    [
        [
            "اورام الثدي - الغده - الكبد",
            "طنطا\nالاسكندرية\nالإسكندرية\nدلالات\nتكلفة عملية\nسعر عملية\nكانسر الثدي\nعلاج ورم حميد\nاورام الثدي الحميدة\nالغدة الدرقية بالليزر",
            "طنطا/الإسكندرية فقط لو مش مطلوب leads من المحافظات دي.",
        ],
        [
            "بنكرياس - قولون - مراره",
            "كيف يتم التبرز\nاستئصال القولون بالكامل\nاستئصال القولون كامل\nاستئصال القولون بالمنظار\nعملية استئصال القولون\nتجربتي مع استئصال جزء من القولون\nاعراض ما بعد عملية استئصال جزء من القولون\nمدة الشفاء من عملية استئصال القولون\nهل استئصال جزء من القولون خطير\nهل عملية استئصال القولون خطيرة\nهل يمكن استئصال القولون بالكامل\nهل ينفع استئصال القولون\nسرطان المعده له علاج\nورم القناة المرارية الحميد\nاعراض ورم البنكرياس الحميد\nأعراض ورم البنكرياس الحميد\nعلاج ورم البنكرياس الحميد\nهل الورم في البنكرياس خطير",
            "تنظيف ضروري قبل أي توسيع أو Smart Bidding.",
        ],
    ],
)

add_heading(doc, "Keywords تتضاف حسب كل Ad Group", 1)
add_table(
    doc,
    ["Ad Group", "Exact Match", "Phrase Match", "قرار"],
    [
        [
            "اورام الكبد",
            "[افضل دكتور اورام في مصر]\n[دكتور اورام كبد بالقاهرة]\n[دكتور اورام كبد]\n[علاج اورام الكبد]\n[علاج سرطان الكبد]\n[جراح اورام كبد]\n[جراحة اورام الكبد]",
            '"افضل دكتور اورام كبد"\n"دكتور اورام كبد"\n"جراحة اورام الكبد"',
            "أقوى ad group. زوّدي هنا بثقة.",
        ],
        [
            "الغدة الدرقيه",
            "[عملية الغدة الدرقية]\n[اورام الغدة الدرقية]\n[استئصال الغدة الدرقية]\n[استئصال ورم الغدة الدرقية]\n[دكتور جراحة الغدة الدرقية]",
            '"عملية الغدة الدرقية"\n"استئصال ورم الغدة الدرقية"\n"دكتور جراحة الغدة الدرقية"',
            "لا تزودي ورم الغدة الدرقية حاليًا لأنه صرف بدون conversions.",
        ],
        [
            "أورام الثدي",
            "[دكتور اورام ثدي]\n[جراحة اورام الثدي]\n[افضل دكتور اورام الثدي في القاهرة]\n[دكتور جراحة اورام الثدي]",
            '"دكتور اورام ثدي"\n"جراحة اورام الثدي"\n"دكتور جراحة اورام الثدي"',
            "قللي علاج سرطان الثدي وعلاج اورام الثدي.",
        ],
        [
            "أورام البنكرياس",
            "[اورام البنكرياس]\n[دكتور اورام بنكرياس]\n[جراحة اورام البنكرياس]",
            "لا تضيفي Phrase دلوقتي",
            "اختبار صغير فقط بسبب conversion واحد.",
        ],
    ],
)

add_heading(doc, "Keywords لا تتضاف دلوقتي", 1)
add_copy_block(
    doc,
    "استبعاد من الإضافة:",
    [
        "علاج كانسر الثدي",
        "علاج ورم حميد في الغدة الدرقية",
        "العلاج الهرموني لسرطان الثدي",
        "دلالات اورام الغدة الدرقية",
        "كيف يتم التبرز بعد استئصال القولون",
        "ورم القناة المرارية الحميد",
        "ورم المعدة الحميد",
        "استئصال القولون",
        "علاج سرطان القولون",
        "علاج سرطان المعدة",
        "ورم البنكرياس",
        "ورم على البنكرياس",
    ],
)

add_heading(doc, "تعديلات الإعلانات حسب Ad Group", 1)
add_table(
    doc,
    ["Ad Group", "تعديلات مطلوبة", "Headlines مقترحة"],
    [
        [
            "اورام الكبد",
            "عدلي: حاصل على ماجستير الجراحةالعامة -> حاصل على ماجستير الجراحة العامة\nعدلي: خبرةفي استئصال أورام الكبد -> خبرة في استئصال أورام الكبد\nقللي: أفضل علاج لأورام الكبد / رعاية متخصصة لحالات الكبد",
            "علاج أورام الكبد\nدكتور أورام كبد بالقاهرة\nجراحة أورام الكبد\nاستئصال ورم من الكبد\nاستشارة متخصصة لأورام الكبد\nد. كيرلس مدحت جراح أورام",
        ],
        [
            "الغدة الدرقيه",
            "عدلي: اسشر دكتور كيرلس مدحت -> استشر دكتور كيرلس مدحت",
            "استئصال الغدة الدرقية\nعملية الغدة الدرقية\nدكتور جراحة الغدة الدرقية\nعلاج أورام الغدة الدرقية\nاستئصال ورم الغدة الدرقية\nاستشر د. كيرلس مدحت",
        ],
        [
            "أورام الثدي",
            "وحّدي اسم الدكتور: كيرلس أو كرولوس حسب الاسم المعتمد.\nقللي العبارات العامة: رعاية متخصصة لكل حالة / خطة علاج مناسبة للحالة / متابعة قبل وبعد العلاج",
            "دكتور أورام ثدي\nجراحة أورام الثدي\nاستشارة متخصصة للثدي\nتقييم أورام الثدي\nاحجزي كشفك الآن\nد. كيرلس مدحت جراح أورام",
        ],
        [
            "المعدة / القولون / المرارة",
            "أوقفي المعدة مؤقتًا.\nأوقفي القولون مؤقتًا.\nأوقفي المرارة مؤقتًا.\nلو هيفضلوا شغالين خليهم جراحة high intent فقط.",
            "لا توسّعي الرسائل العامة. ركزي على جراحة/استئصال/دكتور متخصص.",
        ],
    ],
)

add_heading(doc, "Location", 1)
add_bullets(
    doc,
    [
        "لا تغيّري الاستهداف الجغرافي الأساسي دلوقتي.",
        "الإعداد المطلوب: Egypt.",
        "Location option: Presence only - People in or regularly in targeted locations.",
        "ما تستخدميش Presence or interest.",
        "لو الميزانية ضغطت: زودي أولوية القاهرة والجيزة بعد ظهور city-level data.",
    ],
)

add_heading(doc, "Schedule", 1)
add_table(
    doc,
    ["اليوم", "الجدول المقترح"],
    [
        ["Saturday", "12:00 PM - 12:00 AM"],
        ["Sunday", "12:00 PM - 12:00 AM"],
        ["Monday", "12:00 PM - 12:00 AM"],
        ["Tuesday", "12:00 PM - 12:00 AM"],
        ["Wednesday", "12:00 PM - 12:00 AM"],
        ["Thursday", "12:00 PM - 12:00 AM"],
        ["Friday", "12:00 PM - 12:00 AM"],
    ],
)

add_heading(doc, "Bidding / Max Conversions", 1)
add_table(
    doc,
    ["Campaign", "القرار", "Target CPA", "ملاحظات"],
    [
        ["اورام الثدي - الغده - الكبد", "حوّلي إلى Maximize Conversions", "300 EGP", "لو الصرف قل أو Google رفض: ارفعي لـ 350 EGP. ما تنزليش تحت 280 EGP."],
        ["بنكرياس - قولون - مراره", "لا تحوّلي دلوقتي", "غير مناسب حاليًا", "صرف 2125.90 EGP وجاب conversion واحد فقط. محتاج تنظيف الأول."],
    ],
)

add_heading(doc, "Budget", 1)
add_bullets(
    doc,
    [
        "اورام الثدي - الغده - الكبد: 850 - 900 EGP/day.",
        "بنكرياس - قولون - مراره: 100 - 150 EGP/day.",
        "زودي الكبد، ثبتي الغدة، قللي الثدي شوية لحد ما يتحسن CPA.",
        "قللي أو أوقفي المعدة/القولون/المرارة.",
        "سيبي البنكرياس اختبار صغير فقط.",
    ],
)

add_heading(doc, "ترتيب التنفيذ", 1)
add_bullets(
    doc,
    [
        "ضيفي Account negatives.",
        "ضيفي Campaign negatives.",
        "ضيفي keywords الجديدة Exact/Phrase حسب كل ad group.",
        "عدلي typos في الإعلانات.",
        "وحّدي اسم الدكتور.",
        "أوقفي المعدة/القولون/المرارة أو قلليهم جدًا.",
        "حوّلي كامبين اورام الثدي - الغده - الكبد إلى Maximize Conversions بـ Target CPA 300 EGP.",
        "راجعي بعد 5-7 أيام بدون تعديلات كبيرة في النص.",
    ],
)

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        set_run_font(run, run.font.size.pt if run.font.size else 10.5, run.bold, None)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
